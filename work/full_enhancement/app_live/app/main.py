from __future__ import annotations
from .media_center import router as media_router
from .photo_archive import router as photo_router
from fastapi import FastAPI, UploadFile, File, Form, HTTPException, Request, Response
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from pathlib import Path
from datetime import datetime
import tempfile, shutil, os, json

from .security_hardening import ensure_security_tables,password_policy,is_locked,record_fail,clear_fail,touch_session,idle_expired
from .mail_config_api import router as mail_config_router
from .auth_service import ensure_default_admin,authenticate,create_session,get_user_by_session,delete_session,public_user,has_permission,required_permission,PERMISSION_LABELS,ROLE_DEFAULTS,ALL_PERMISSIONS,hash_password,permissions_json
from .auth_context import current_user_id,current_username,current_ip
from .db import init_db, connect, now, log, backup_db, BASE
from .calc import sale_calc, stock_calc, get_settings
from .xlsx_service import seed_from_master_if_empty, import_stock_source, import_sales_source, export_master, master_preview, apply_master_import
from .pdf_service import sales_report, stock_report, purchase_report, consultant_report, expertise_report, profit_loss_report, critical_stock_report, management_report

app=FastAPI(title="Çayan Tarsus RENEW PRO",version="8.0.0")
app.include_router(photo_router)
app.include_router(mail_config_router)
app.include_router(media_router)
STATIC=BASE/"app"/"static"
app.mount("/static",StaticFiles(directory=STATIC),name="static")

@app.middleware("http")
async def no_cache(request,call_next):
    r=await call_next(request)
    r.headers["Cache-Control"]="no-store, no-cache, must-revalidate, max-age=0"
    r.headers["Pragma"]="no-cache";r.headers["Expires"]="0"
    return r



@app.middleware("http")
async def auth_middleware(request:Request,call_next):
    path=request.url.path
    # Static UI and login API remain public.
    if path.startswith("/static/") or path=="/" or path in ("/api/auth/login","/api/build-version","/api/version"):
        return await call_next(request)
    if path.startswith("/api/"):
        token=request.cookies.get("renew_session")
        user=get_user_by_session(token)
        if not user:
            return JSONResponse({"detail":"Oturum gerekli."},status_code=401)
        if idle_expired(user.get("session_last_seen")):
            delete_session(token)
            return JSONResponse({"detail":"30 dakika işlem yapılmadığı için oturum kapatıldı."},status_code=401)
        if request.method in ("POST","PUT","PATCH","DELETE") and path not in ("/api/auth/logout",):
            csrf=request.headers.get("x-csrf-token","")
            if not csrf or csrf!=user.get("session_csrf"):
                return JSONResponse({"detail":"Güvenlik doğrulaması başarısız (CSRF)."},status_code=403)
        touch_session(token)
        t1=current_user_id.set(user["id"]);t2=current_username.set(user["username"]);t3=current_ip.set(request.client.host if request.client else "")
        request.state.user=user
        try:
            perm=required_permission(path,request.method)
            if perm and not has_permission(user,perm):
                return JSONResponse({"detail":"Bu işlem için yetkiniz yok.","permission":perm},status_code=403)
            return await call_next(request)
        finally:
            current_user_id.reset(t1);current_username.reset(t2);current_ip.reset(t3)
    return await call_next(request)

class LoginIn(BaseModel):
    username:str
    password:str

class PasswordIn(BaseModel):
    old_password:str=""
    new_password:str

class UserIn(BaseModel):
    username:str
    full_name:str
    password:str=""
    role:str="viewer"
    permissions:list[str]=[]
    active:bool=True
    must_change_password:bool=False

@app.post("/api/auth/login")
def auth_login(x:LoginIn,response:Response,request:Request):
    locked,until=is_locked(x.username)
    if locked: raise HTTPException(429,f"Hesap {until.strftime('%H:%M')} saatine kadar geçici kilitli.")
    user=authenticate(x.username,x.password)
    if not user:
        count,lock=record_fail(x.username)
        if lock: raise HTTPException(429,"5 hatalı giriş nedeniyle hesap 15 dakika kilitlendi.")
        raise HTTPException(401,f"Kullanıcı adı veya şifre hatalı. Kalan deneme: {5-count}")
    clear_fail(x.username)
    token,csrf=create_session(user["id"],request.client.host if request.client else "",request.headers.get("user-agent","")[:250])
    response.set_cookie("renew_session",token,httponly=True,secure=True,samesite="strict",max_age=43200)
    response.set_cookie("renew_csrf",csrf,httponly=False,secure=True,samesite="strict",max_age=43200)
    d=public_user(user);d["csrf_token"]=csrf
    return d

@app.post("/api/auth/logout")
def auth_logout(request:Request,response:Response):
    user=request.state.user
    log("auth",user["id"],"LOGOUT","Sistemden çıkış yaptı")
    delete_session(request.cookies.get("renew_session"))
    response.delete_cookie("renew_session")
    return {"ok":True}

@app.get("/api/auth/me")
def auth_me(request:Request):
    return public_user(request.state.user)

@app.post("/api/auth/change-password")
def change_password(x:PasswordIn,request:Request):
    u=request.state.user
    password_policy(x.new_password)
    if u["must_change_password"] or verify_password_for_user(u,x.old_password):
        with connect() as con:con.execute("UPDATE users SET password_hash=?,must_change_password=0,updated_at=? WHERE id=?",(hash_password(x.new_password),now(),u["id"]))
        log("user",u["id"],"PASSWORD_CHANGE","Kendi şifresini değiştirdi")
        return {"ok":True}
    raise HTTPException(400,"Mevcut şifre yanlış.")

def verify_password_for_user(u,password):
    from .auth_service import verify_password
    return verify_password(password,u["password_hash"])

@app.get("/api/users")
def users_list():
    with connect() as con:
        rows=[dict(r) for r in con.execute("SELECT * FROM users ORDER BY role='administrator' DESC,full_name").fetchall()]
    return [public_user(r) for r in rows]

@app.get("/api/users/meta")
def users_meta():
    groups={
      "Görüntüleme":["dashboard.view","stocks.view","sales.view","acquisitions.view","expertise.view","performance.view","settings.view","reports.view","media.view","audit.view"],
      "Veri Düzenleme":["stocks.edit","sales.edit","acquisitions.edit","expertise.edit","media.edit"],
      "Sistem İşlemleri":["settings.edit","imports.use","backup.use"],
      "Yönetim":["users.manage"]
    }
    return {"permissions":PERMISSION_LABELS,"roles":{k:v for k,v in ROLE_DEFAULTS.items()},"groups":groups}

@app.post("/api/users")
def user_create(x:UserIn):
    if not x.username.strip() or not x.full_name.strip():raise HTTPException(400,"Kullanıcı adı ve ad soyad zorunlu.")
    password_policy(x.password)
    role=x.role if x.role in ROLE_DEFAULTS else "viewer"
    perms=ALL_PERMISSIONS if role=="administrator" else x.permissions
    try:
        with connect() as con:
            cur=con.execute("""INSERT INTO users(username,full_name,password_hash,role,permissions,active,must_change_password,created_at,updated_at)
              VALUES(?,?,?,?,?,?,?,?,?)""",(x.username.strip(),x.full_name.strip(),hash_password(x.password),role,permissions_json(role,perms),1 if x.active else 0,1 if x.must_change_password else 0,now(),now()))
            uid=cur.lastrowid
        log("user",uid,"CREATE",f"Kullanıcı oluşturuldu: {x.username}")
        return {"ok":True,"id":uid}
    except Exception as e:
        if "UNIQUE" in str(e).upper():raise HTTPException(400,"Bu kullanıcı adı zaten kullanılıyor.")
        raise

@app.put("/api/users/{uid}")
def user_update(uid:int,x:UserIn):
    with connect() as con:
        old=con.execute("SELECT * FROM users WHERE id=?",(uid,)).fetchone()
        if not old:raise HTTPException(404,"Kullanıcı bulunamadı.")
        role=x.role if x.role in ROLE_DEFAULTS else "viewer"
        if old["role"]=="administrator" and (role!="administrator" or not x.active):
            n=con.execute("SELECT COUNT(*) n FROM users WHERE role='administrator' AND active=1 AND id<>?",(uid,)).fetchone()["n"]
            if n<1:raise HTTPException(400,"Son aktif administrator hesabı pasif yapılamaz veya rolü değiştirilemez.")
        perms=ALL_PERMISSIONS if role=="administrator" else x.permissions
        fields=[x.username.strip(),x.full_name.strip(),role,permissions_json(role,perms),1 if x.active else 0,1 if x.must_change_password else 0,now()]
        sql="UPDATE users SET username=?,full_name=?,role=?,permissions=?,active=?,must_change_password=?,updated_at=?"
        if x.password:
            password_policy(x.password)
            sql+=",password_hash=?";fields.append(hash_password(x.password))
        sql+=" WHERE id=?";fields.append(uid)
        con.execute(sql,tuple(fields))
    log("user",uid,"UPDATE",f"Kullanıcı yetkileri güncellendi: {x.username}",json.dumps(dict(old),ensure_ascii=False,default=str),json.dumps({"username":x.username,"role":role,"permissions":perms,"active":x.active},ensure_ascii=False))
    return {"ok":True}

@app.delete("/api/users/{uid}")
def user_delete(uid:int,request:Request):
    if uid==request.state.user["id"]:raise HTTPException(400,"Kendi kullanıcı hesabınızı silemezsiniz.")
    with connect() as con:
        u=con.execute("SELECT * FROM users WHERE id=?",(uid,)).fetchone()
        if not u:raise HTTPException(404,"Kullanıcı bulunamadı.")
        if u["role"]=="administrator":
            n=con.execute("SELECT COUNT(*) n FROM users WHERE role='administrator' AND active=1").fetchone()["n"]
            if n<=1:raise HTTPException(400,"Son aktif administrator silinemez.")
        con.execute("DELETE FROM users WHERE id=?",(uid,))
    log("user",uid,"DELETE",f"Kullanıcı silindi: {u['username']}")
    return {"ok":True}



@app.get("/api/network-info")
def network_info(request:Request):
    import socket
    ip="127.0.0.1"
    s=socket.socket(socket.AF_INET,socket.SOCK_DGRAM)
    try:
        s.connect(("8.8.8.8",80));ip=s.getsockname()[0]
    except Exception:
        try:ip=socket.gethostbyname(socket.gethostname())
        except Exception:pass
    finally:
        try:s.close()
        except:pass
    return {"lan_ip":ip,"port":8765,"localhost":"http://127.0.0.1:8765","lan_url":f"http://{ip}:8765"}

@app.get("/api/build-version")
def build_version():
    return {"version":"8.0.0","monthly_acquisitions":True,"dashboard_redesign":True,"profile":True}

@app.on_event("startup")
def startup():
    init_db();ensure_security_tables();ensure_default_admin();seed_from_master_if_empty()

@app.get("/",response_class=HTMLResponse)
def home(): return (STATIC/"index.html").read_text(encoding="utf-8")

@app.get("/api/version")
def version(): return {"version":"8.0.0","name":"Çayan Tarsus RENEW PRO"}


def _save_upload_temp(upload_bytes: bytes, suffix: str) -> Path:
    """Windows-safe temporary upload file. mkstemp fd is explicitly closed before writing."""
    fd, name = tempfile.mkstemp(suffix=suffix)
    os.close(fd)
    path = Path(name)
    path.write_bytes(upload_bytes)
    return path

class StockIn(BaseModel):
    plate:str; model_year:int|None=None; km:int=0; brand:str=""; model:str=""; version:str=""; color:str=""
    fuel:str=""; transmission:str=""; purchase_date:str=""; purchase_type:str="NAKİT"
    purchase_price:float=0; list_price:float=0; target_profit:float=0

class AcquisitionIn(BaseModel):
    purchase_date:str; plate:str; model_year:int|None=None; km:int=0; brand:str=""; model:str=""; version:str=""
    color:str=""; fuel:str=""; transmission:str=""; purchase_type:str="NAKİT"; acquired_by:str
    purchase_price:float=0; expense:float=0

class SaleIn(BaseModel):
    plate:str; vehicle_info:str=""; model_year:int|None=None; purchase_type:str="NAKİT"; purchase_price:float=0
    purchase_date:str=""; sale_date:str; extra_expense:float=0; customer:str=""; sale_type:str="NAKİT"
    sale_price:float=0; consultant:str=""; insurance_income:float=0; credit_income:float=0; warranty_income:float=0


def _norm_choice(v):
    import unicodedata
    s=unicodedata.normalize("NFKD",str(v or "").strip().upper())
    s="".join(ch for ch in s if not unicodedata.combining(ch))
    return " ".join(s.replace("İ","I").split())

def _require_purchase_type(v):
    with connect() as con:
        rows=[r["name"] for r in con.execute("SELECT name FROM purchase_rules WHERE active=1").fetchall()]
    if not v or not any(_norm_choice(v)==_norm_choice(x) for x in rows):
        raise HTTPException(400,f"Alım Türü zorunlu ve Parametreler'deki aktif seçeneklerden biri olmalı: {v!r}")

def _require_sale_fields(v):
    _require_purchase_type(v.purchase_type)
    with connect() as con:
        sale_types=[r["name"] for r in con.execute("SELECT name FROM sale_types WHERE active=1").fetchall()]
        consultants=[r["name"] for r in con.execute("SELECT name FROM consultants WHERE active=1").fetchall()]
    if not v.sale_type or not any(_norm_choice(v.sale_type)==_norm_choice(x) for x in sale_types):
        raise HTTPException(400,f"Satış Türü zorunlu ve aktif seçeneklerden biri olmalı: {v.sale_type!r}")
    if not v.consultant or not any(_norm_choice(v.consultant)==_norm_choice(x) for x in consultants):
        raise HTTPException(400,f"Satış Danışmanı zorunlu ve aktif danışmanlardan biri olmalı: {v.consultant!r}")



@app.get("/api/vehicle-sources")
def vehicle_sources():
    with connect() as con:
        stocks=[dict(r) for r in con.execute("SELECT plate,brand,model,version,model_year FROM vehicles WHERE status='STOCK' ORDER BY brand,model,plate").fetchall()]
        sales=[dict(r) for r in con.execute("SELECT plate,vehicle_info,model_year,sale_date FROM sales ORDER BY sale_date DESC,id DESC").fetchall()]
        acquisitions=[dict(r) for r in con.execute("SELECT plate,brand,model,version,model_year,purchase_date,acquired_by FROM acquisitions ORDER BY purchase_date DESC,id DESC").fetchall()]
    return {"stocks":stocks,"sales":sales,"acquisitions":acquisitions}

@app.get("/api/vehicle-card/{plate}")
def vehicle_card(plate:str):
    import re as _re
    pkey=_re.sub(r"[^A-Za-z0-9_-]","",str(plate or "").upper())
    if not pkey: raise HTTPException(400,"Plaka zorunlu.")
    with connect() as con:
        stock=con.execute("""SELECT * FROM vehicles WHERE upper(replace(plate,' ',''))=upper(replace(?,' ','')) ORDER BY id DESC LIMIT 1""",(plate,)).fetchone()
        sale=con.execute("""SELECT * FROM sales WHERE upper(replace(plate,' ',''))=upper(replace(?,' ','')) ORDER BY sale_date DESC,id DESC LIMIT 1""",(plate,)).fetchone()
        acq=con.execute("""SELECT * FROM acquisitions WHERE upper(replace(plate,' ',''))=upper(replace(?,' ','')) ORDER BY purchase_date DESC,id DESC LIMIT 1""",(plate,)).fetchone()
    sr=dict(stock) if stock else None; sl=dict(sale) if sale else None; aq=dict(acq) if acq else None
    status="SATILDI" if sl else ("STOKTA" if sr and sr.get("status")=="STOCK" else (sr.get("status") if sr else "KAYIT"))
    calc_data=sale_calc(sl) if sl else (stock_calc(sr) if sr else None)
    media_root=BASE/"vehicle_media"/pkey
    def media_count(kind):
        d=media_root/kind
        if not d.exists(): return 0
        return sum(1 for f in d.iterdir() if f.is_file() and f.suffix.lower() in {".jpg",".jpeg",".png",".webp"})
    return {"plate":plate,"status":status,"stock":sr,"sale":sl,"acquisition":aq,"calc":calc_data,
            "photos":{"original":media_count("original"),"processed":media_count("processed")},
            "vehicle_info":((f"{sr.get('brand','')} {sr.get('model','')} {sr.get('version','')}".strip() if sr else "") or (sl.get("vehicle_info","") if sl else ""))}

@app.get("/api/dashboard")
def dashboard(month:str):
    def prev_month(m):
        y,mo=map(int,m.split("-"));mo-=1
        if mo==0:y-=1;mo=12
        return f"{y:04d}-{mo:02d}"
    pm=prev_month(month)
    with connect() as con:
        stocks=[dict(r) for r in con.execute("SELECT * FROM vehicles WHERE status='STOCK' ORDER BY purchase_date,id").fetchall()]
        sales=[dict(r) for r in con.execute("SELECT * FROM sales WHERE substr(sale_date,1,7)=? ORDER BY sale_date,id",(month,)).fetchall()]
        prev=[dict(r) for r in con.execute("SELECT * FROM sales WHERE substr(sale_date,1,7)=?",(pm,)).fetchall()]
        purchases=[dict(r) for r in con.execute("SELECT * FROM vehicles WHERE substr(purchase_date,1,7)=?",(month,)).fetchall()]
        exp=con.execute("SELECT COALESCE(SUM(done_count),0)d,COALESCE(SUM(converted_count),0)c FROM expertise WHERE month=?",(month,)).fetchone()
        recent=[dict(r) for r in con.execute("SELECT * FROM audit_log ORDER BY id DESC LIMIT 10").fetchall()]
        trend_sales=[dict(r) for r in con.execute("SELECT * FROM sales WHERE sale_date<=? ORDER BY sale_date,id",(month+"-31",)).fetchall()]
        acquisition_rows=[dict(r) for r in con.execute("SELECT * FROM acquisitions WHERE substr(purchase_date,1,7)=?",(month,)).fetchall()]
    sc=[stock_calc(x) for x in stocks];cc=[sale_calc(x) for x in sales];pc=[sale_calc(x) for x in prev]
    rev=sum(x["sale_price"] for x in sales);perf=sum(x["performance_profit"] for x in cc);prevrev=sum(x["sale_price"] for x in prev);prevperf=sum(x["performance_profit"] for x in pc)
    consultants={}
    for r,c in zip(sales,cc):
        n=(r.get("consultant") or "BELİRTİLMEMİŞ").strip() or "BELİRTİLMEMİŞ"
        d=consultants.setdefault(n,{"name":n,"sales_count":0,"revenue":0,"profit":0,"sks":0})
        d["sales_count"]+=1;d["revenue"]+=r["sale_price"];d["profit"]+=c["performance_profit"];d["sks"]+=c["sks_days"]
    cs=[]
    for d in consultants.values():
        d["avg_sks"]=d["sks"]/d["sales_count"] if d["sales_count"] else 0;d.pop("sks");cs.append(d)
    cs.sort(key=lambda x:x["profit"],reverse=True)
    brands={}
    for r in stocks:brands[r["brand"] or "DİĞER"]=brands.get(r["brand"] or "DİĞER",0)+1
    pts={}
    for r in sales:pts[r["purchase_type"] or "DİĞER"]=pts.get(r["purchase_type"] or "DİĞER",0)+1
    pairs=[{"plate":r["plate"],"vehicle_info":r["vehicle_info"],"consultant":r["consultant"],"sks":c["sks_days"],"profit":c["performance_profit"]} for r,c in zip(sales,cc)]
    crit=[{"plate":r["plate"],"vehicle":f"{r['brand']} {r['model']} {r['version']}".strip(),"sks":c["sks_days"],"purchase_price":r["purchase_price"],"list_price":r["list_price"],"finance":c["sks_finance"],"total_cost":c["total_cost"],"profit":c["estimated_profit"]} for r,c in zip(stocks,sc)]
    months=[];y,mo=map(int,month.split("-"))
    for _ in range(12):
        months.append(f"{y:04d}-{mo:02d}");mo-=1
        if mo==0:y-=1;mo=12
    months.reverse();trend=[]
    for tm in months:
        rows=[r for r in trend_sales if str(r.get("sale_date") or "")[:7]==tm]
        calcs=[sale_calc(r) for r in rows]
        trend.append({"month":tm,"sales_count":len(rows),"revenue":sum(float(r.get("sale_price") or 0) for r in rows),"profit":sum(float(c.get("performance_profit") or 0) for c in calcs)})
    source_map={}
    for r in acquisition_rows:
        key=(r.get("acquired_by") or "BELİRTİLMEMİŞ").strip() or "BELİRTİLMEMİŞ"
        d=source_map.setdefault(key,{"name":key,"count":0,"value":0});d["count"]+=1;d["value"]+=float(r.get("purchase_price") or 0)
    alerts=[]
    critical_count=sum(c["sks_days"]>90 for c in sc);loss_risk=sum(c["estimated_profit"]<0 for c in sc)
    missing_info=sum(not str(r.get("brand") or "").strip() or not str(r.get("model") or "").strip() or not float(r.get("list_price") or 0) for r in stocks)
    media_missing=0
    for r in stocks:
        pkey="".join(ch for ch in str(r.get("plate") or "").upper() if ch.isalnum())
        folder=BASE/"vehicle_media"/pkey/"original"
        if not folder.exists() or not any(x.is_file() and x.suffix.lower() in {".jpg",".jpeg",".png",".webp"} for x in folder.iterdir()):media_missing+=1
    if critical_count:alerts.append({"level":"critical","title":"90+ gün kritik stok","count":critical_count,"target":"stocks","filter":"90_plus"})
    if loss_risk:alerts.append({"level":"warning","title":"Zarar riski bulunan stok","count":loss_risk,"target":"stocks","filter":"loss"})
    if missing_info:alerts.append({"level":"info","title":"Eksik araç bilgisi","count":missing_info,"target":"stocks","filter":"missing"})
    if media_missing:alerts.append({"level":"info","title":"Fotoğrafı eksik stok","count":media_missing,"target":"photos","filter":"missing_media"})
    return {
      "stock_count":len(stocks),"stock_purchase_total":sum(r["purchase_price"] for r in stocks),"stock_cost_total":sum(c["total_cost"] for c in sc),
      "stock_sale_value":sum(r["list_price"] for r in stocks),"stock_estimated_profit":sum(c["estimated_profit"] for c in sc),"stock_sks_finance":sum(c["sks_finance"] for c in sc),
      "stock_avg_sks":sum(c["sks_days"] for c in sc)/len(sc) if sc else 0,"stock_critical_count":sum(c["sks_days"]>90 for c in sc),
      "sales_count":len(sales),"sales_revenue":rev,"sales_net_profit":sum(c["net_profit"] for c in cc),"sales_performance_profit":perf,
      "avg_sale_profit":perf/len(cc) if cc else 0,"avg_sale_sks":sum(c["sks_days"] for c in cc)/len(cc) if cc else 0,"avg_margin":perf/rev if rev else 0,
      "profitable_sales":sum(c["performance_profit"]>=0 for c in cc),"losing_sales":sum(c["performance_profit"]<0 for c in cc),
      "previous_sales_count":len(prev),"previous_revenue":prevrev,"previous_performance_profit":prevperf,
      "sales_count_change":(len(sales)-len(prev))/len(prev) if prev else (1 if sales else 0),
      "revenue_change":(rev-prevrev)/prevrev if prevrev else (1 if rev else 0),"profit_change":(perf-prevperf)/abs(prevperf) if prevperf else (1 if perf else 0),
      "purchased_count":len(purchases),"purchased_value":sum(r["purchase_price"] for r in purchases),
      "expertise_done":exp["d"],"expertise_converted":exp["c"],"expertise_rate":exp["c"]/exp["d"] if exp["d"] else 0,
      "stock_risk":{"0_30":sum(c["sks_days"]<=30 for c in sc),"31_60":sum(31<=c["sks_days"]<=60 for c in sc),"61_90":sum(61<=c["sks_days"]<=90 for c in sc),"90_plus":sum(c["sks_days"]>90 for c in sc)},
      "consultant_summary":cs,"brand_distribution":[{"name":k,"count":v} for k,v in sorted(brands.items(),key=lambda z:z[1],reverse=True)],
      "purchase_type_distribution":[{"name":k,"count":v} for k,v in sorted(pts.items(),key=lambda z:z[1],reverse=True)],
      "top_sales":sorted(pairs,key=lambda x:x["profit"],reverse=True)[:5],"loss_sales":sorted([x for x in pairs if x["profit"]<0],key=lambda x:x["profit"])[:5],
      "critical_stock":sorted(crit,key=lambda x:x["sks"],reverse=True)[:10],"recent_activity":recent,
      "monthly_trend":trend,"management_alerts":alerts,
      "acquisition_sources":sorted(source_map.values(),key=lambda x:(x["count"],x["value"]),reverse=True),
      "stock_capital":sum(c["total_cost"] for c in sc),"stock_loss_risk_count":loss_risk,"stock_missing_info_count":missing_info,"stock_missing_media_count":media_missing
    }

@app.get("/api/profile")
def profile(request:Request):
    user=request.state.user
    with connect() as con:
        actions=[dict(r) for r in con.execute("SELECT * FROM audit_log WHERE username=? ORDER BY id DESC LIMIT 100",(user["username"],)).fetchall()]
    by_entity={}
    for row in actions:
        key=row.get("entity") or "other";by_entity[key]=by_entity.get(key,0)+1
    return {"user":public_user(user),"summary":{"total_actions":len(actions),"by_entity":by_entity,"last_action":actions[0]["created_at"] if actions else None},"actions":actions}

@app.get("/api/stocks")
def stocks():
    with connect() as con:rows=[dict(r) for r in con.execute("SELECT * FROM vehicles WHERE status='STOCK' ORDER BY purchase_date,id").fetchall()]
    return [{**r,"calc":stock_calc(r)} for r in rows]

@app.post("/api/stocks")
def add_stock(x:StockIn):
    _require_purchase_type(x.purchase_type)
    vals=x.model_dump();t=now()
    with connect() as con:
        ex=con.execute("SELECT id FROM vehicles WHERE upper(replace(plate,' ',''))=upper(replace(?,' ',''))",(x.plate,)).fetchone()
        if ex:
            con.execute("UPDATE vehicles SET "+",".join(f"{k}=?" for k in vals)+",status='STOCK',updated_at=? WHERE id=?",tuple(vals.values())+(t,ex["id"]));vid=ex["id"];act="UPDATE"
        else:
            keys=list(vals)+["status","created_at","updated_at"];cur=con.execute(f"INSERT INTO vehicles({','.join(keys)}) VALUES({','.join('?'*len(keys))})",tuple(vals.values())+("STOCK",t,t));vid=cur.lastrowid;act="CREATE"
    log("vehicle",vid,act,x.plate);return {"ok":True,"id":vid}

@app.put("/api/stocks/{vid}")
def update_stock(vid:int,x:StockIn):
    _require_purchase_type(x.purchase_type)
    vals=x.model_dump()
    with connect() as con:
        con.execute("UPDATE vehicles SET "+",".join(f"{k}=?" for k in vals)+",updated_at=? WHERE id=?",tuple(vals.values())+(now(),vid))
    log("vehicle",vid,"UPDATE",x.plate);return {"ok":True}

@app.delete("/api/stocks/{vid}")
def delete_stock(vid:int):
    with connect() as con:r=con.execute("SELECT plate FROM vehicles WHERE id=?",(vid,)).fetchone();con.execute("DELETE FROM vehicles WHERE id=?",(vid,))
    log("vehicle",vid,"DELETE",r["plate"] if r else "");return {"ok":True}

@app.post("/api/stocks/{vid}/sell")
def sell_stock(vid:int,x:SaleIn):
    _require_sale_fields(x)
    vals=x.model_dump();t=now()
    with connect() as con:
        v=con.execute("SELECT * FROM vehicles WHERE id=?",(vid,)).fetchone()
        if not v:raise HTTPException(404,"Stok aracı bulunamadı")
        vals["vehicle_id"]=vid;keys=list(vals)+["created_at","updated_at"]
        cur=con.execute(f"INSERT INTO sales({','.join(keys)}) VALUES({','.join('?'*len(keys))})",tuple(vals.values())+(t,t))
        con.execute("UPDATE vehicles SET status='SOLD',updated_at=? WHERE id=?",(t,vid))
    log("vehicle",vid,"SELL",x.plate);log("sale",cur.lastrowid,"CREATE",x.plate);return {"ok":True}

@app.get("/api/sales")
def sales(month:str|None=None):
    with connect() as con:
        rows=[dict(r) for r in (con.execute("SELECT * FROM sales WHERE substr(sale_date,1,7)=? ORDER BY sale_date,id",(month,)).fetchall() if month else con.execute("SELECT * FROM sales ORDER BY sale_date,id").fetchall())]
    return [{**r,"calc":sale_calc(r)} for r in rows]

@app.post("/api/sales")
def add_sale(x:SaleIn):
    _require_sale_fields(x)
    vals=x.model_dump();t=now()
    with connect() as con:
        ex=con.execute("""SELECT id FROM sales WHERE upper(replace(plate,' ',''))=upper(replace(?,' ','')) AND sale_date=? ORDER BY id LIMIT 1""",(x.plate,x.sale_date)).fetchone()
        if ex:
            con.execute("UPDATE sales SET "+",".join(f"{k}=?" for k in vals)+",updated_at=? WHERE id=?",tuple(vals.values())+(t,ex["id"]))
            sid=ex["id"]; action="UPDATE"
        else:
            keys=list(vals)+["created_at","updated_at"]
            cur=con.execute(f"INSERT INTO sales({','.join(keys)}) VALUES({','.join('?'*len(keys))})",tuple(vals.values())+(t,t))
            sid=cur.lastrowid; action="CREATE"
    log("sale",sid,action,x.plate);return {"ok":True,"id":sid}

@app.put("/api/sales/{sid}")
def update_sale(sid:int,x:SaleIn):
    _require_sale_fields(x)
    vals=x.model_dump()
    with connect() as con:con.execute("UPDATE sales SET "+",".join(f"{k}=?" for k in vals)+",updated_at=? WHERE id=?",tuple(vals.values())+(now(),sid))
    log("sale",sid,"UPDATE",x.plate);return {"ok":True}

@app.delete("/api/sales/{sid}")
def delete_sale(sid:int):
    with connect() as con:r=con.execute("SELECT plate FROM sales WHERE id=?",(sid,)).fetchone();con.execute("DELETE FROM sales WHERE id=?",(sid,))
    log("sale",sid,"DELETE",r["plate"] if r else "");return {"ok":True}


def _require_acquisition_person(name):
    with connect() as con:
        people={_norm_choice(r["consultant"]) for r in con.execute("SELECT DISTINCT consultant FROM expertise").fetchall()}
        people |= {_norm_choice(r["name"]) for r in con.execute("SELECT name FROM consultants WHERE active=1").fetchall()}
    if not name or _norm_choice(name) not in people:
        raise HTTPException(400,f"Alınan kişi / ekspertiz personeli zorunlu ve sistemde kayıtlı olmalı: {name!r}")

def _sync_acquisition_expertise(month,person):
    if not month or not person:return
    with connect() as con:
        n=con.execute("SELECT COUNT(*) n FROM acquisitions WHERE substr(purchase_date,1,7)=? AND upper(acquired_by)=upper(?)",(month,person)).fetchone()["n"]
        old=con.execute("SELECT done_count FROM expertise WHERE consultant=? AND month=?",(person,month)).fetchone()
        done=old["done_count"] if old else 0
        con.execute("""INSERT INTO expertise(consultant,month,done_count,converted_count,updated_at) VALUES(?,?,?,?,?)
        ON CONFLICT(consultant,month) DO UPDATE SET converted_count=excluded.converted_count,updated_at=excluded.updated_at""",(person,month,done,n,now()))

@app.get("/api/acquisitions")
def acquisitions(month:str):
    with connect() as con:return [dict(r) for r in con.execute("SELECT * FROM acquisitions WHERE substr(purchase_date,1,7)=? ORDER BY purchase_date,id",(month,)).fetchall()]

@app.get("/api/acquisitions/summary")
def acquisitions_summary(month:str):
    with connect() as con:rows=[dict(r) for r in con.execute("SELECT * FROM acquisitions WHERE substr(purchase_date,1,7)=? ORDER BY purchase_date,id",(month,)).fetchall()]
    bp={};bt={}
    for r in rows:
        d=bp.setdefault(r["acquired_by"],{"name":r["acquired_by"],"count":0,"purchase_total":0,"expense_total":0,"total_cost":0})
        d["count"]+=1;d["purchase_total"]+=r["purchase_price"];d["expense_total"]+=r["expense"];d["total_cost"]+=r["purchase_price"]+r["expense"]
        bt[r["purchase_type"]]=bt.get(r["purchase_type"],0)+1
    return {"count":len(rows),"purchase_total":sum(r["purchase_price"] for r in rows),"expense_total":sum(r["expense"] for r in rows),"total_cost":sum(r["purchase_price"]+r["expense"] for r in rows),"by_person":sorted(bp.values(),key=lambda x:(-x["count"],x["name"])),"by_type":[{"name":k,"count":v} for k,v in bt.items()]}

@app.post("/api/acquisitions")
def add_acquisition(x:AcquisitionIn):
    _require_purchase_type(x.purchase_type);_require_acquisition_person(x.acquired_by);vals=x.model_dump();t=now()
    with connect() as con:
        keys=list(vals);cur=con.execute(f"INSERT INTO acquisitions({','.join(keys)},created_at,updated_at) VALUES({','.join('?'*len(keys))},?,?)",tuple(vals.values())+(t,t));aid=cur.lastrowid
    _sync_acquisition_expertise(x.purchase_date[:7],x.acquired_by);log("acquisition",aid,"CREATE",f"{x.plate} • {x.acquired_by}");return {"ok":True,"id":aid}

@app.put("/api/acquisitions/{aid}")
def update_acquisition(aid:int,x:AcquisitionIn):
    _require_purchase_type(x.purchase_type);_require_acquisition_person(x.acquired_by)
    with connect() as con:
        old=con.execute("SELECT purchase_date,acquired_by FROM acquisitions WHERE id=?",(aid,)).fetchone();vals=x.model_dump()
        con.execute("UPDATE acquisitions SET "+",".join(f"{k}=?" for k in vals)+",updated_at=? WHERE id=?",tuple(vals.values())+(now(),aid))
    if old:_sync_acquisition_expertise(old["purchase_date"][:7],old["acquired_by"])
    _sync_acquisition_expertise(x.purchase_date[:7],x.acquired_by);log("acquisition",aid,"UPDATE",x.plate);return {"ok":True}

@app.delete("/api/acquisitions/{aid}")
def delete_acquisition(aid:int):
    with connect() as con:
        old=con.execute("SELECT * FROM acquisitions WHERE id=?",(aid,)).fetchone();con.execute("DELETE FROM acquisitions WHERE id=?",(aid,))
    if old:_sync_acquisition_expertise(old["purchase_date"][:7],old["acquired_by"])
    log("acquisition",aid,"DELETE",old["plate"] if old else "");return {"ok":True}



@app.post("/api/import/acquisitions/preview")
async def preview_acquisitions(month:str=Form(...), file:UploadFile=File(...)):
    ext=Path(file.filename or "").suffix.lower()
    if ext not in (".xlsx",".xlsm",".xls"):
        raise HTTPException(400,"Aylık alım Excel dosyası .xlsx/.xlsm/.xls olmalı.")
    from openpyxl import load_workbook
    tmp=BASE/"imports"/f"acq_{datetime.now():%Y%m%d_%H%M%S_%f}{ext}"
    tmp.parent.mkdir(exist_ok=True)
    tmp.write_bytes(await file.read())
    try:
        wb=load_workbook(tmp,data_only=False,read_only=True)
        ws=wb["AYLIK ALIM TAKİBİ"] if "AYLIK ALIM TAKİBİ" in wb.sheetnames else wb[wb.sheetnames[0]]
        rows=[]
        for rn in range(5,ws.max_row+1):
            plate=str(ws.cell(rn,4).value or "").strip()
            if not plate: continue
            def num(v):
                try:return float(v or 0)
                except:return 0
            d=ws.cell(rn,2).value
            if hasattr(d,"strftime"): d=d.strftime("%Y-%m-%d")
            else: d=str(d or "").strip()
            rows.append({
              "purchase_date":d,"plate":plate,"model_year":int(num(ws.cell(rn,5).value)) if ws.cell(rn,5).value else None,
              "km":int(num(ws.cell(rn,6).value)),"brand":str(ws.cell(rn,7).value or "").strip(),
              "model":str(ws.cell(rn,8).value or "").strip(),"version":str(ws.cell(rn,9).value or "").strip(),
              "color":str(ws.cell(rn,10).value or "").strip(),"fuel":str(ws.cell(rn,11).value or "").strip(),
              "transmission":str(ws.cell(rn,12).value or "").strip(),"purchase_type":str(ws.cell(rn,13).value or "").strip(),
              "acquired_by":str(ws.cell(rn,14).value or "").strip(),"purchase_price":num(ws.cell(rn,15).value),
              "expense":num(ws.cell(rn,16).value)
            })
        return {"rows":rows,"count":len(rows),"file_name":file.filename,"month":month}
    except Exception as e:
        raise HTTPException(400,f"Aylık alım Excel okunamadı: {e}")
    finally:
        tmp.unlink(missing_ok=True)

@app.post("/api/import/acquisitions/apply")
def apply_acquisitions(payload:dict):
    rows=payload.get("rows",[])
    month=payload.get("month","")
    mode=payload.get("mode","merge")
    if not rows: raise HTTPException(400,"Aktarılacak aylık alım kaydı yok.")
    if mode not in ("merge","replace_month"): raise HTTPException(400,"Geçersiz aktarım modu.")
    backup_db()
    if mode=="replace_month" and month:
        with connect() as con: con.execute("DELETE FROM acquisitions WHERE substr(purchase_date,1,7)=?",(month,))
    count=0
    for r in rows:
        item=AcquisitionIn(**r)
        _require_purchase_type(item.purchase_type); _require_acquisition_person(item.acquired_by)
        # same plate+date update, else insert
        with connect() as con:
            ex=con.execute("SELECT id FROM acquisitions WHERE upper(replace(plate,' ',''))=upper(replace(?,' ','')) AND purchase_date=? ORDER BY id LIMIT 1",(item.plate,item.purchase_date)).fetchone()
        if ex: update_acquisition(ex["id"],item)
        else: add_acquisition(item)
        count+=1
    return {"ok":True,"count":count,"mode":mode,"month":month}

@app.get("/api/expertise")
def expertise(month:str):
    # Personel listesi EXPERTİZ TAKİP kayıtlarının tamamından gelir.
    # Seçili ayda kayıt yoksa kişi yine görünür ve değerleri 0 olur.
    with connect() as con:
        names=[r["consultant"] for r in con.execute("SELECT DISTINCT consultant FROM expertise ORDER BY consultant").fetchall()]
        selected={r["consultant"]:dict(r) for r in con.execute("SELECT * FROM expertise WHERE month=?",(month,)).fetchall()}
    return [{
        "consultant":name,"month":month,
        "done_count":int(selected.get(name,{}).get("done_count",0) or 0),
        "converted_count":int(selected.get(name,{}).get("converted_count",0) or 0)
    } for name in names]

@app.post("/api/expertise")
def save_expertise(payload:dict):
    m=payload["month"]
    with connect() as con:
        for r in payload["rows"]:
            con.execute("""INSERT INTO expertise(consultant,month,done_count,converted_count,updated_at) VALUES(?,?,?,?,?)
            ON CONFLICT(consultant,month) DO UPDATE SET done_count=excluded.done_count,converted_count=excluded.converted_count,updated_at=excluded.updated_at""",(r["consultant"],m,int(r.get("done_count",0)),int(r.get("converted_count",0)),now()))
    log("expertise",m,"UPDATE",f"{len(payload['rows'])} personel");return {"ok":True}

@app.post("/api/expertise/person")
def add_expertise_person(payload:dict):
    name=(payload.get("name") or "").strip();m=payload.get("month")
    if not name or not m:raise HTTPException(400,"Personel ve ay zorunlu")
    with connect() as con:con.execute("INSERT OR IGNORE INTO expertise(consultant,month,done_count,converted_count,updated_at) VALUES(?,?,0,0,?)",(name,m,now()))
    return {"ok":True}

@app.delete("/api/expertise/person")
def remove_expertise_person(name:str,month:str):
    with connect() as con:con.execute("DELETE FROM expertise WHERE consultant=? AND month=?",(name,month))
    return {"ok":True}

@app.get("/api/performance")
def performance(month:str):
    with connect() as con:
        cons=[dict(r) for r in con.execute("SELECT * FROM consultants WHERE active=1 ORDER BY name").fetchall()]
        sales=[dict(r) for r in con.execute("SELECT * FROM sales WHERE substr(sale_date,1,7)=?",(month,)).fetchall()]
    out=[]
    for c in cons:
        rr=[r for r in sales if (r.get("consultant") or "").strip().upper()==c["name"].strip().upper()];cc=[sale_calc(r) for r in rr];perf=sum(x["performance_profit"] for x in cc)
        out.append({**c,"sales_count":len(rr),"revenue":sum(r["sale_price"] for r in rr),"performance_profit":perf,"avg_profit":perf/len(rr) if rr else 0,"avg_sks":sum(x["sks_days"] for x in cc)/len(cc) if cc else 0,"sales_target_rate":len(rr)/c["target_sales"] if c["target_sales"] else 0})
    return out

@app.get("/api/settings")
def settings():
    with connect() as con:
        return {"settings":get_settings(),"rules":[dict(r) for r in con.execute("SELECT * FROM purchase_rules ORDER BY rowid").fetchall()],
                "sale_types":[dict(r) for r in con.execute("SELECT * FROM sale_types ORDER BY rowid").fetchall()],
                "income_types":[dict(r) for r in con.execute("SELECT * FROM income_types ORDER BY rowid").fetchall()],
                "consultants":[dict(r) for r in con.execute("SELECT * FROM consultants ORDER BY id").fetchall()]}

@app.post("/api/settings")
def save_settings(payload:dict):
    with connect() as con:
        for k,v in payload.get("settings",{}).items():
            con.execute("""INSERT INTO settings(key,value,value_type,updated_at) VALUES(?,?,'number',?)
            ON CONFLICT(key) DO UPDATE SET value=excluded.value,updated_at=excluded.updated_at""",(k,str(v),now()))
        for r in payload.get("rules",[]):
            con.execute("""INSERT INTO purchase_rules(name,apply_sks,apply_fixed,active) VALUES(?,?,?,?)
            ON CONFLICT(name) DO UPDATE SET apply_sks=excluded.apply_sks,apply_fixed=excluded.apply_fixed,active=excluded.active""",(r["name"],int(r["apply_sks"]),int(r["apply_fixed"]),int(r["active"])))
        con.execute("DELETE FROM sale_types")
        for r in payload.get("sale_types",[]): con.execute("INSERT INTO sale_types(name,active) VALUES(?,?)",(r["name"],int(r["active"])))
        con.execute("DELETE FROM income_types")
        for r in payload.get("income_types",[]): con.execute("INSERT INTO income_types(name,active,description) VALUES(?,?,?)",(r["name"],int(r["active"]),r.get("description","")))
        # consultant upserts, do not blindly delete people referenced by sales
        for r in payload.get("consultants",[]):
            con.execute("""INSERT INTO consultants(name,active,target_sales,target_profit) VALUES(?,?,?,?)
            ON CONFLICT(name) DO UPDATE SET active=excluded.active,target_sales=excluded.target_sales,target_profit=excluded.target_profit""",(r["name"],int(r["active"]),int(r.get("target_sales",0)),float(r.get("target_profit",0))))
    log("settings","all","UPDATE","Parametreler güncellendi");return {"ok":True}

@app.post("/api/import/stock/preview")
async def preview_stock(file:UploadFile=File(...)):
    ext=Path(file.filename or "").suffix.lower()
    if ext not in (".xlsx",".xlsm"):
        raise HTTPException(400,"Stok aktarımı için .xlsx veya .xlsm dosyası yükleyin. Eski .xls biçimi desteklenmiyor.")
    tmp=_save_upload_temp(await file.read(),ext)
    try:
        rows=import_stock_source(tmp)
        return {"rows":rows,"file_name":file.filename,"count":len(rows)}
    except Exception as e:
        raise HTTPException(400,f"Stok Excel okunamadı: {e}")
    finally:tmp.unlink(missing_ok=True)

def _norm_choice(v):
    import unicodedata
    s=str(v or "").strip().upper().replace("İ","I").replace("Ş","S").replace("Ğ","G").replace("Ü","U").replace("Ö","O").replace("Ç","C")
    return " ".join(s.split())

def _canonical_choice(value, choices):
    n=_norm_choice(value)
    for c in choices:
        if _norm_choice(c)==n: return c
    return None

def _active_choices(table):
    with connect() as con:
        return [r["name"] for r in con.execute(f"SELECT name FROM {table} WHERE active=1 ORDER BY rowid").fetchall()]

@app.post("/api/import/stock/apply")
def apply_stock(payload:dict):
    rows=payload.get("rows",[])
    mode=payload.get("mode","merge")
    if mode not in ("merge","replace"): raise HTTPException(400,"Geçersiz stok aktarım modu")
    backup_db()
    if mode=="replace":
        with connect() as con:
            con.execute("DELETE FROM vehicles WHERE status='STOCK'")
    purchase_types=_active_choices("purchase_rules")
    for i,r in enumerate(rows,1):
        c=_canonical_choice(r.get("purchase_type"),purchase_types)
        if not c: raise HTTPException(400,f"{i}. stok kaydı ({r.get('plate','')}): Alım Türü zorunlu ve Parametreler listesinden seçilmelidir.")
        r["purchase_type"]=c
    count=0
    for r in rows:
        add_stock(StockIn(**r)); count+=1
    log("stock_import","all",mode.upper(),f"{count} araç")
    return {"ok":True,"count":count,"mode":mode}

@app.post("/api/import/sales/preview")
async def preview_sales(month:str=Form(...),file:UploadFile=File(...)):
    ext=Path(file.filename or "").suffix.lower()
    if ext not in (".xlsx",".xlsm"):
        raise HTTPException(400,"Satış aktarımı için .xlsx veya .xlsm dosyası yükleyin. Eski .xls biçimi desteklenmiyor.")
    tmp=_save_upload_temp(await file.read(),ext)
    try:
        rows=import_sales_source(tmp,month)
        return {"rows":rows,"file_name":file.filename,"month":month,"count":len(rows)}
    except Exception as e:
        raise HTTPException(400,f"Satış Excel okunamadı: {e}")
    finally:tmp.unlink(missing_ok=True)

@app.post("/api/import/sales/apply")
def apply_sales(payload:dict):
    rows=payload.get("rows",[])
    mode=payload.get("mode","merge")
    month=payload.get("month","")
    if mode not in ("merge","replace_month"): raise HTTPException(400,"Geçersiz satış aktarım modu")
    purchase_types=_active_choices("purchase_rules"); sale_types=_active_choices("sale_types"); consultants=_active_choices("consultants")
    for i,r in enumerate(rows,1):
        pt=_canonical_choice(r.get("purchase_type"),purchase_types); st=_canonical_choice(r.get("sale_type"),sale_types); co=_canonical_choice(r.get("consultant"),consultants)
        missing=[]
        if not pt: missing.append("Alım Türü")
        if not st: missing.append("Satış Türü")
        if not co: missing.append("Danışman")
        if missing: raise HTTPException(400,f"{i}. satış kaydı ({r.get('plate','')}): {', '.join(missing)} seçimi zorunlu.")
        r["purchase_type"]=pt; r["sale_type"]=st; r["consultant"]=co
    backup_db()
    if mode=="replace_month":
        if not month: raise HTTPException(400,"Ay seçimi zorunlu")
        with connect() as con:
            con.execute("DELETE FROM sales WHERE substr(sale_date,1,7)=?",(month,))
    count=0
    for r in rows:
        # Merge modunda aynı plaka + satış tarihi varsa güncelle, tekrar ekleme.
        with connect() as con:
            ex=con.execute("SELECT id FROM sales WHERE upper(replace(plate,' ',''))=upper(replace(?,' ','')) AND sale_date=?",(r.get("plate",""),r.get("sale_date",""))).fetchone()
        if ex:
            update_sale(ex["id"],SaleIn(**r))
        else:
            add_sale(SaleIn(**r))
        count+=1
    log("sales_import",month or "all",mode.upper(),f"{count} satış")
    return {"ok":True,"count":count,"mode":mode,"month":month}

@app.post("/api/master/preview")
async def master_preview_api(file:UploadFile=File(...)):
    ext=Path(file.filename or "").suffix.lower()
    if ext not in (".xlsx",".xlsm"):
        raise HTTPException(400,"Ana yedek .xlsx veya .xlsm olmalıdır. HTML dosyası Excel yedeği olarak yüklenemez.")
    tmp=_save_upload_temp(await file.read(),ext)
    try:
        d=master_preview(tmp)
        token=f"master_{datetime.now():%Y%m%d_%H%M%S_%f}{ext}"
        dest=BASE/"imports"/token; dest.parent.mkdir(parents=True,exist_ok=True); shutil.copy2(tmp,dest)
        return {"token":token,"file_name":file.filename,"summary":d["summary"],"sample_sales":d["sales"][:5],"sample_stocks":d["stocks"][:5]}
    except Exception as e:
        raise HTTPException(400,f"Ana Excel okunamadı: {e}")
    finally:
        tmp.unlink(missing_ok=True)

@app.post("/api/master/apply")
def master_apply(payload:dict):
    p=BASE/"imports"/Path(payload["token"]).name
    if not p.exists():raise HTTPException(404,"Önizleme dosyası yok")
    return {"ok":True,"summary":apply_master_import(p,payload.get("mode","replace"),payload.get("file_name","RENEW.xlsx"))}

@app.get("/api/import-history")
def import_history():
    with connect() as con:return [dict(r) for r in con.execute("SELECT * FROM import_history ORDER BY id DESC LIMIT 50").fetchall()]

@app.get("/api/export/excel")
def export_excel(month:str|None=None):
    try:
        (BASE/"reports").mkdir(parents=True,exist_ok=True)
        out=BASE/"reports"/f"RENEW_GUNCEL_{datetime.now():%Y%m%d_%H%M%S_%f}.xlsx"
        export_master(out,month)
        if not out.exists() or out.stat().st_size<1000:
            raise RuntimeError("Excel çıktı dosyası oluşmadı")
        return FileResponse(out,filename=f"RENEW_GUNCEL_{datetime.now():%Y%m%d_%H%M}.xlsx",media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    except Exception as e:
        raise HTTPException(400,f"Excel dışa aktarım hatası: {e}")
@app.get("/api/report/sales")
def rs(month:str):p=sales_report(month);return FileResponse(p,filename=p.name)
@app.get("/api/report/stock")
def rst():p=stock_report();return FileResponse(p,filename=p.name)
@app.get("/api/report/purchases")
def rp(month:str):p=purchase_report(month);return FileResponse(p,filename=p.name)
@app.get("/api/report/consultants")
def rc(month:str):p=consultant_report(month);return FileResponse(p,filename=p.name)
@app.get("/api/report/expertise")
def rexp(month:str):p=expertise_report(month);return FileResponse(p,filename=p.name)
@app.get("/api/report/profit-loss")
def rpl(month:str):p=profit_loss_report(month);return FileResponse(p,filename=p.name)
@app.get("/api/report/critical-stock")
def rcs():p=critical_stock_report();return FileResponse(p,filename=p.name)
@app.get("/api/report/management")
def rm(month:str):p=management_report(month);return FileResponse(p,filename=p.name)

@app.get("/api/report/acquisitions")
def racq(month:str):
    from .pdf_service import acquisition_report
    p=acquisition_report(month);return FileResponse(p,filename=p.name)

@app.get("/api/export/acquisitions.xlsx")
def xacq(month:str):
    from .acquisition_xlsx import export_acquisitions
    p=export_acquisitions(month);return FileResponse(p,filename=p.name,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.get("/api/template/acquisitions.xlsx")
def tacq():
    p=BASE/"templates"/"AYLIK_ALIM_TAKIP_SABLONU.xlsx"
    return FileResponse(p,filename="RENEW_Aylik_Arac_Alim_Takip_Sablonu.xlsx",media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.post("/api/backup")
def backup():
    p=backup_db();return {"ok":bool(p),"file":p.name if p else None}

@app.get("/api/backups")
def backups():
    root=BASE/"backups";root.mkdir(parents=True,exist_ok=True)
    rows=[]
    for p in sorted(root.glob("renew_*.db"),key=lambda x:x.stat().st_mtime,reverse=True)[:100]:
        st=p.stat();rows.append({"name":p.name,"size":st.st_size,"created_at":datetime.fromtimestamp(st.st_mtime).isoformat(timespec="seconds")})
    return rows

@app.get("/api/backups/{name}")
def download_backup(name:str):
    clean=Path(name).name
    if clean!=name or not clean.startswith("renew_") or not clean.endswith(".db"):raise HTTPException(400,"Geçersiz yedek adı")
    p=BASE/"backups"/clean
    if not p.is_file():raise HTTPException(404,"Yedek bulunamadı")
    return FileResponse(p,filename=clean,media_type="application/octet-stream")
@app.get("/api/audit")
def audit(username:str="",entity:str="",q:str="",limit:int=500):
    sql="SELECT * FROM audit_log WHERE 1=1";args=[]
    if username:sql+=" AND username=?";args.append(username)
    if entity:sql+=" AND entity=?";args.append(entity)
    if q:sql+=" AND (details LIKE ? OR entity_id LIKE ? OR old_value LIKE ? OR new_value LIKE ?)";args += [f"%{q}%"]*4
    sql+=" ORDER BY id DESC LIMIT ?";args.append(min(max(limit,1),2000))
    with connect() as con:return [dict(r) for r in con.execute(sql,args).fetchall()]
